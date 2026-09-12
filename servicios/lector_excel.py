"""Lectura del Excel de trabajadores y patrón.

Se adapta a la estructura real de la base del despacho:
- Encabezados en dos niveles: una fila de grupos en celdas combinadas
  (p. ej. DOMICILIO o SALARIO DIARIO IMSS) y, debajo, la fila de encabezados.
- Filas vacías o con texto auxiliar del formato (p. ej. "FECHA DE INGRESO"
  repetido) que no deben contarse como registros.

Este módulo solo lee y depura; no decide qué falta (eso es del validador).
"""
from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime, time
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel


class ErrorLectura(Exception):
    """El libro no se puede procesar (dañado, sin hoja de trabajadores o sin encabezados)."""


# Catálogo de campos conocidos: clave interna → (nombre legible, encabezados posibles).
# Los encabezados están normalizados (sin acentos ni signos, en mayúsculas) y pueden
# incluir el grupo de la fila superior: "DOMICILIO NUMERO" ≠ "SALARIO DIARIO IMSS NUMERO".
CAMPOS_TRABAJADOR = {
    "nombre": ("Nombre", ["NOMBRE", "NOMBRE COMPLETO", "NOMBRE DEL TRABAJADOR"]),
    "nacionalidad": ("Nacionalidad", ["NACIONALIDAD"]),
    "lugar_nacimiento": ("Lugar de nacimiento", ["LUGAR DE NACIMIENTO", "LUGAR NACIMIENTO", "ORIGINARIO DE"]),
    "fecha_nacimiento": ("Fecha de nacimiento", ["FECHA DE NACIMIENTO", "FECHA NACIMIENTO"]),
    "sexo": ("Sexo", ["SEXO"]),
    "estado_civil": ("Estado civil", ["ESTADO CIVIL"]),
    "edad": ("Edad", ["EDAD"]),
    "domicilio_numero": ("Domicilio: número", ["DOMICILIO NUMERO"]),
    "domicilio_calle": ("Domicilio: calle", ["DOMICILIO CALLE", "CALLE"]),
    "domicilio_colonia": ("Domicilio: colonia", ["DOMICILIO COLONIA", "COLONIA"]),
    "domicilio_municipio": ("Domicilio: municipio", ["DOMICILIO MUNICIPIO", "MUNICIPIO"]),
    "credencial_elector": ("Credencial de elector",
                           ["CRED ELECTOR", "CREDENCIAL DE ELECTOR", "CREDENCIAL ELECTOR", "INE"]),
    "seguro_social": ("Seguro social", ["SEGURO SOCIAL", "NSS", "NUMERO DE SEGURO SOCIAL"]),
    "fecha_ingreso": ("Fecha de ingreso", ["FECHA DE INGRESO"]),
    "curp": ("CURP", ["CURP"]),
    "rfc": ("RFC", ["RFC"]),
    "beneficiarios": ("Beneficiarios", ["BENEFICIARIOS"]),
    "correo": ("Correo electrónico", ["CORREO ELECTRONICO", "CORREO", "EMAIL"]),
    "puesto": ("Puesto", ["PUESTO"]),
    "jornada": ("Jornada de trabajo (horario)",
                ["JORNADA DE TRABAJO HORARIO", "JORNADA DE TRABAJO", "JORNADA", "HORARIO"]),
    "horario_comida": ("Horario de comida", ["HORARIO DE COMIDA"]),
    "lv_entrada": ("Entrada de lunes a viernes",
                   ["HORARIO LUNES A VIERNES ENTRADA", "HORA DE ENTRADA", "ENTRADA"]),
    "lv_salida": ("Salida de lunes a viernes",
                  ["HORARIO LUNES A VIERNES SALIDA", "HORA DE SALIDA", "SALIDA"]),
    "sabado_entrada": ("Entrada del sábado", ["HORARIO SABADO ENTRADA", "SABADO ENTRADA", "SABADO HORA DE ENTRADA"]),
    "sabado_salida": ("Salida del sábado", ["HORARIO SABADO SALIDA", "SABADO SALIDA", "SABADO HORA DE SALIDA"]),
    "salario_numero": ("Salario diario IMSS (número)",
                       ["SALARIO DIARIO IMSS NUMERO", "SALARIO DIARIO IMSS"]),
    "salario_letra": ("Salario diario IMSS (letra)", ["SALARIO DIARIO IMSS LETRA"]),
    "dia_pago": ("Día de pago", ["DIA DE PAGO"]),
    "dia_descanso": ("Día de descanso", ["DIA DE DESCANSO"]),
    "fecha_alta_imss": ("Fecha de alta IMSS", ["FECHA ALTA IMSS", "FECHA DE ALTA IMSS"]),
    "vigencia": ("Vigencia del contrato", ["VIGENCIA DEL CONTRATO", "VIGENCIA"]),
}

CAMPOS_PATRON = {
    "razon_social": ("Razón social", ["RAZON SOCIAL", "NOMBRE O RAZON SOCIAL", "DENOMINACION"]),
    "actividad": ("Actividad de la empresa", ["ACTIVIDAD DE LA EMPRESA", "ACTIVIDAD", "GIRO"]),
    "rfc": ("RFC", ["RFC"]),
    "domicilio_numero": ("Domicilio: número", ["DOMICILIO NUMERO"]),
    "domicilio_calle": ("Domicilio: calle", ["DOMICILIO CALLE", "CALLE"]),
    "domicilio_colonia": ("Domicilio: colonia", ["DOMICILIO COLONIA", "COLONIA"]),
    "domicilio_municipio": ("Domicilio: municipio", ["DOMICILIO MUNICIPIO", "MUNICIPIO"]),
    "domicilio_cp": ("Código postal", ["DOMICILIO CODIGO POSTAL", "CODIGO POSTAL", "CP"]),
    "representante_legal": ("Representante legal", ["REPRESENTANTE LEGAL"]),
    "acta_numero": ("Acta constitutiva: número", ["ACTA CONSTITUTIVA NUMERO"]),
    "acta_fecha": ("Acta constitutiva: fecha", ["ACTA CONSTITUTIVA FECHA"]),
    "notario_nombre": ("Nombre del notario",
                       ["ACTA CONSTITUTIVA NOMBRE NOTARIO", "NOMBRE NOTARIO", "NOMBRE DEL NOTARIO"]),
    "notario_numero": ("Número del notario",
                       ["ACTA CONSTITUTIVA NUMERO NOTARIO", "NUMERO NOTARIO", "NUMERO DE NOTARIO"]),
    "acta_ciudad": ("Ciudad", ["ACTA CONSTITUTIVA CIUDAD", "CIUDAD"]),
}

# Campos que sí pueden contener fechas; en los demás, una fecha es un error de formato de la celda.
CAMPOS_FECHA = {"fecha_ingreso", "fecha_alta_imss", "fecha_nacimiento", "vigencia", "acta_fecha"}

# Campos que por sí solos bastan para considerar que una fila es un registro real.
CAMPOS_IDENTIDAD = {"nombre", "curp", "rfc", "seguro_social", "razon_social"}

PALABRAS_INSTRUCCION = ("NOTA", "INSTRUCC", "LLENAR", "FAVOR DE", "EJEMPLO", "IMPORTANTE", "OBSERVAC")
FILAS_BUSQUEDA_ENCABEZADOS = 15
MAX_COLUMNAS = 200
MAX_FILAS_VACIAS_SEGUIDAS = 500  # evita recorrer hojas con formato aplicado hasta la fila 1 048 576


def normalizar(texto) -> str:
    """Texto sin acentos ni signos, en mayúsculas, para comparar encabezados."""
    texto = unicodedata.normalize("NFKD", str(texto))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"[^A-Za-z0-9]+", " ", texto).strip().upper()


def a_texto(valor) -> str:
    """Convierte el contenido de una celda en texto, sin alterar el dato."""
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    if isinstance(valor, datetime):
        formato = "%d/%m/%Y" if valor.time() == time(0) else "%d/%m/%Y %H:%M"
        return valor.strftime(formato)
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, time):
        return valor.strftime("%H:%M")
    if isinstance(valor, float):
        return str(int(valor)) if valor.is_integer() else ("%.6f" % valor).rstrip("0")
    return str(valor).strip()


def leer_libro(contenido: bytes) -> dict:
    """Lee las hojas de trabajadores y patrón; las demás se reportan como omitidas."""
    try:
        libro = load_workbook(BytesIO(contenido), data_only=True)
    except Exception as error:  # ZIP dañado o contenido que no es un libro de Excel
        raise ErrorLectura("No fue posible abrir el libro de Excel; puede estar dañado.") from error

    hoja_trabajadores = _buscar_hoja(libro, "TRABAJADORES", "TRABAJADOR")
    if hoja_trabajadores is None:
        nombres = ", ".join(f"«{h.title}»" for h in libro.worksheets)
        raise ErrorLectura(f"No se encontró la hoja «Trabajadores». Hojas del libro: {nombres}.")
    hoja_patron = _buscar_hoja(libro, "PATRON", "EMPRESA")
    if hoja_patron is hoja_trabajadores:
        hoja_patron = None

    trabajadores = _leer_hoja(hoja_trabajadores, CAMPOS_TRABAJADOR)
    if trabajadores is None:
        raise ErrorLectura(f"No se localizaron los encabezados en la hoja «{hoja_trabajadores.title}».")

    avisos = []
    patron = None
    if hoja_patron is None:
        avisos.append("El libro no tiene una hoja de patrón; los datos patronales quedan pendientes.")
    else:
        patron = _leer_hoja(hoja_patron, CAMPOS_PATRON)
        if patron is None:
            avisos.append(f"No se localizaron los encabezados en la hoja «{hoja_patron.title}».")

    leidas = [h for h in (hoja_trabajadores, hoja_patron) if h is not None]
    return {
        "hojas_leidas": [h.title for h in leidas],
        "hojas_omitidas": [
            {"nombre": h.title, "vacia": _hoja_vacia(h)} for h in libro.worksheets if h not in leidas
        ],
        "trabajadores": trabajadores,
        "patron": patron,
        "avisos": avisos,
    }


def _buscar_hoja(libro, nombre_exacto: str, *fragmentos: str):
    for hoja in libro.worksheets:
        if normalizar(hoja.title) == nombre_exacto:
            return hoja
    for hoja in libro.worksheets:
        if any(fragmento in normalizar(hoja.title) for fragmento in fragmentos):
            return hoja
    return None


def _hoja_vacia(hoja) -> bool:
    return all(
        a_texto(valor) == ""
        for fila in hoja.iter_rows(max_row=min(hoja.max_row, 1000), values_only=True)
        for valor in fila
    )


def _leer_hoja(hoja, catalogo: dict) -> dict | None:
    """Detecta los encabezados (uno o dos niveles) y separa registros de filas auxiliares."""
    alias = {a: clave for clave, (_, lista) in catalogo.items() for a in lista}
    max_col = min(hoja.max_column, MAX_COLUMNAS)

    # Las celdas combinadas solo guardan el valor en su esquina superior izquierda;
    # en los encabezados se replica a todas las columnas que abarcan.
    combinadas = {}
    for rango in hoja.merged_cells.ranges:
        valor = hoja.cell(rango.min_row, rango.min_col).value
        for f in range(rango.min_row, min(rango.max_row, FILAS_BUSQUEDA_ENCABEZADOS + 1) + 1):
            for c in range(rango.min_col, min(rango.max_col, max_col) + 1):
                combinadas[(f, c)] = valor

    def encabezado(f: int, c: int) -> str:
        if f < 1:
            return ""
        return a_texto(combinadas.get((f, c), hoja.cell(f, c).value))

    def identificar(grupo: str, texto: str) -> str | None:
        completo = normalizar(f"{grupo} {texto}")
        return alias.get(completo) or alias.get(normalizar(texto))

    # 1. La fila de encabezados es la que reconoce más campos del catálogo.
    fila_enc, mejor = None, 0
    for f in range(1, min(hoja.max_row, FILAS_BUSQUEDA_ENCABEZADOS) + 1):
        puntaje = sum(
            1 for c in range(1, max_col + 1)
            if encabezado(f, c) and identificar(encabezado(f - 1, c), encabezado(f, c))
        )
        if puntaje > mejor:
            fila_enc, mejor = f, puntaje
    if fila_enc is None or mejor < 2:
        return None

    # 2. Columnas: encabezado real + grupo de la fila superior (si existe).
    columnas, usadas, duplicadas = [], set(), []
    usa_fila_superior = False
    for c in range(1, max_col + 1):
        texto, grupo = encabezado(fila_enc, c), encabezado(fila_enc - 1, c)
        usa_fila_superior = usa_fila_superior or bool(grupo)
        if not texto and grupo:  # encabezado escrito solo en la fila superior
            texto, grupo = grupo, ""
        if not texto:
            continue
        letra = get_column_letter(c)
        clave = identificar(grupo, texto)
        if clave in usadas:
            duplicadas.append(f"{letra} ({texto})")
            clave = None
        usadas.add(clave)
        columnas.append({
            "indice": c,
            "letra": letra,
            "grupo": grupo or None,
            "encabezado": texto,
            "etiqueta": f"{grupo} / {texto}" if grupo else texto,
            "clave": clave or f"columna_{letra}",
            "reconocida": clave is not None,
        })

    # Textos que, si aparecen en una fila de datos, son parte del formato y no un dato.
    textos_formato = set()
    for col in columnas:
        textos_formato.update({normalizar(col["encabezado"]), normalizar(col["etiqueta"])})
        if col["grupo"]:
            textos_formato.add(normalizar(col["grupo"]))

    # 3. Filas de datos.
    indices = {col["indice"] for col in columnas}
    registros, ignoradas, sin_encabezado = [], [], set()
    vacias_seguidas = 0
    for f in range(fila_enc + 1, hoja.max_row + 1):
        valores, auxiliares = {}, {}
        for col in columnas:
            valor = hoja.cell(f, col["indice"]).value
            if isinstance(valor, date) and col["reconocida"] and col["clave"] not in CAMPOS_FECHA:
                # Celda con formato de fecha en un campo que no es fecha: Excel muestra
                # EDAD 29 como 29/01/1900. Se recupera el número que realmente se capturó.
                valor = to_excel(valor)
            texto = a_texto(valor)
            if not texto:
                continue
            destino = auxiliares if normalizar(texto) in textos_formato else valores
            destino[col["clave"]] = texto
        for c in range(1, max_col + 1):
            if c not in indices and a_texto(hoja.cell(f, c).value):
                sin_encabezado.add(get_column_letter(c))

        if not valores and not auxiliares:
            vacias_seguidas += 1
            if vacias_seguidas >= MAX_FILAS_VACIAS_SEGUIDAS:
                break
            continue
        vacias_seguidas = 0

        motivo = _motivo_para_ignorar(valores, auxiliares)
        if motivo:
            ignoradas.append({"fila": f, "motivo": motivo})
        else:
            registros.append({"fila": f, "valores": valores, "auxiliares": auxiliares})

    return {
        "hoja": hoja.title,
        "filas_encabezado": [fila_enc - 1, fila_enc] if usa_fila_superior else [fila_enc],
        "columnas": columnas,
        "registros": registros,
        "ignoradas": ignoradas,
        "columnas_sin_encabezado": sorted(sin_encabezado, key=lambda l: (len(l), l)),
        "columnas_duplicadas": duplicadas,
        "no_encontrados": [nombre for clave, (nombre, _) in catalogo.items() if clave not in usadas],
    }


def _motivo_para_ignorar(valores: dict, auxiliares: dict) -> str | None:
    """Devuelve por qué una fila no es un registro, o None si sí lo es."""
    if not valores:
        if len(auxiliares) > 1:
            return "encabezado repetido"
        return f"texto del formato «{next(iter(auxiliares.values()))}» sin datos"
    if len(valores) == 1:
        clave, texto = next(iter(valores.items()))
        muestra = texto if len(texto) <= 50 else texto[:47] + "…"
        if normalizar(texto).startswith(PALABRAS_INSTRUCCION) or len(texto) > 80:
            return f"nota o instrucción («{muestra}»)"
        if clave not in CAMPOS_IDENTIDAD:
            return f"dato suelto («{muestra}») sin nombre ni identificación"
    return None
