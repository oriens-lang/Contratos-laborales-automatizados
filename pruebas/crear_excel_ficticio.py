"""Crea «pruebas/Base de prueba FICTICIA.xlsx» con la misma estructura que la base del
despacho y datos INVENTADOS (ningún dato real), para probar lectura, validación y
generación.

Uso:  .venv/bin/python pruebas/crear_excel_ficticio.py
"""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

DESTINO = Path(__file__).with_name("Base de prueba FICTICIA.xlsx")

ENCABEZADOS_TRABAJADOR = [
    "NOMBRE", "NACIONALIDAD", "ESTADO CIVIL", "EDAD", "NÚMERO", "CALLE", "COLONIA", "MUNICIPIO",
    "CRED. ELECTOR", "SEGURO SOCIAL", "FECHA DE INGRESO", "CURP", "RFC", "BENEFICIARIOS",
    "CORREO ELECTRONICO", "PUESTO", "JORNADA DE TRABAJO (Horario)", "HORARIO DE COMIDA", "NÚMERO",
    "LETRA", "DIA DE PAGO", "DIA DE DESCANSO", "FECHA ALTA IMSS", "VIGENCIA DEL CONTRATO",
]
ENCABEZADOS_PATRON = [
    "Razon Social", "Actividad de la empresa", "RFC", "Numero", "Calle", "Colonia", "Municipio",
    "Codigo Postal", "Representante legal", "Numero", "Fecha", "Nombre notario", "Numero notario", "Ciudad",
]


def llenar(hoja, fila: int, datos: dict) -> None:
    for columna, valor in datos.items():
        hoja[f"{columna}{fila}"] = valor


libro = Workbook()

t = libro.active
t.title = "Trabajadores"
t["E1"], t["S1"] = "DOMICILIO", "SALARIO DIARIO IMSS"
t.merge_cells("E1:H1")
t.merge_cells("S1:T1")
for i, texto in enumerate(ENCABEZADOS_TRABAJADOR, 1):
    t.cell(2, i, texto)

# Fila 3: trabajador completo. EDAD y SALARIO con formato de fecha, como en la plantilla real.
llenar(t, 3, dict(
    A="TRABAJADOR DE PRUEBA UNO", B="MEXICANA", C="SOLTERO", D=29, E="123", F="AV. EJEMPLO",
    G="CENTRO", H="GUADALAJARA", I="1234567890123", J="12345678901", K=datetime(2026, 1, 15),
    L="PRUA970101HJCRBN05", M="PRUA970101AB1", N="PERSONA BENEFICIARIA DE PRUEBA 100%",
    O="uno@ejemplo.com", P="AUXILIAR ADMINISTRATIVO", Q="LUNES A VIERNES 9:00 A 18:00",
    R="14:00 A 15:00", S=315.04, T="TRESCIENTOS QUINCE PESOS 04/100 M.N.", U="VIERNES",
    V="DOMINGO", W=datetime(2026, 1, 15), X="TIEMPO INDETERMINADO",
))
t["D3"].number_format = t["S3"].number_format = "mm-dd-yy"

# Filas 4 a 8: texto del formato "FECHA DE INGRESO", como en la plantilla real.
for fila in range(4, 9):
    t[f"K{fila}"] = "FECHA DE INGRESO"
# Fila 4: trabajadora incompleta (K conserva el texto del formato; NSS de 10 dígitos).
llenar(t, 4, dict(A="TRABAJADORA DE PRUEBA DOS", B="MEXICANA", C="CASADA", D=35, F="CALLE FICTICIA",
                  H="ZAPOPAN", J=1234567890, P="VENDEDORA", S=300, U="SÁBADO"))
# Fila 9: encabezado repetido.
for i, texto in enumerate(ENCABEZADOS_TRABAJADOR, 1):
    t.cell(9, i, texto)
# Fila 11: instrucción.
t["A11"] = "NOTA: capturar un trabajador por renglón"
# Fila 12: CURP mal formada, correo inválido y RFC repetido con la fila 3.
llenar(t, 12, dict(A="TRABAJADOR DE PRUEBA TRES", B="MEXICANA", L="ABC123", M="PRUA970101AB1",
                   O="correo-sin-arroba", P="CHOFER"))
# Fila 14: dato suelto sin nombre.
t["B14"] = "MEXICANA"

p = libro.create_sheet("Patron")
p["D1"], p["J1"] = "DOMICILIO", "ACTA CONSTITUTIVA"
p.merge_cells("D1:H1")
p.merge_cells("J1:N1")
for i, texto in enumerate(ENCABEZADOS_PATRON, 1):
    p.cell(2, i, texto)
# La fecha del acta (K) queda vacía a propósito.
llenar(p, 3, dict(A="EMPRESA DE PRUEBA, S.A. DE C.V.", B="COMERCIALIZACIÓN DE PRODUCTOS",
                  C="EPR010101AB1", D="100", E="CALLE FICTICIA", F="COLONIA EJEMPLO", G="ZAPOPAN",
                  H=45000, I="REPRESENTANTE DE PRUEBA", J="12,345", L="NOTARIO DE PRUEBA", M=99,
                  N="GUADALAJARA"))

libro.create_sheet("Hoja3")
libro.save(DESTINO)
print(f"Creado: {DESTINO}")
