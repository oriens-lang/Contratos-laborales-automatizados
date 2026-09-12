"""Crea «plantillas/Formato de captura de trabajadores.xlsx»: el Excel que se envía al cliente.

Mismas hojas que la base del despacho (Trabajadores y Patron) más los datos que pide
la plantilla del contrato: lugar y fecha de nacimiento, sexo y horarios de entrada y
salida. Cada columna ya tiene su formato (texto, fecha, hora o número) para que Excel
no convierta, por ejemplo, la edad en una fecha.

Uso:  .venv/bin/python plantillas/crear_formato_captura.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

DESTINO = Path(__file__).with_name("Formato de captura de trabajadores.xlsx")
FILAS_DE_CAPTURA = 100
TEXTO, FECHA, HORA, ENTERO, IMPORTE = "@", "dd/mm/yyyy", "hh:mm", "0", "#,##0.00"
FUENTE = "Aptos Narrow"

# (grupo, encabezado, formato, ancho)
TRABAJADORES = [
    (None, "NOMBRE", TEXTO, 34),
    (None, "NACIONALIDAD", TEXTO, 14),
    (None, "LUGAR DE NACIMIENTO", TEXTO, 24),
    (None, "FECHA DE NACIMIENTO", FECHA, 14),
    (None, "SEXO", TEXTO, 11),
    (None, "ESTADO CIVIL", TEXTO, 13),
    (None, "EDAD", ENTERO, 7),
    ("DOMICILIO", "NÚMERO", TEXTO, 9),
    ("DOMICILIO", "CALLE", TEXTO, 22),
    ("DOMICILIO", "COLONIA", TEXTO, 20),
    ("DOMICILIO", "MUNICIPIO", TEXTO, 20),
    (None, "CRED. ELECTOR", TEXTO, 20),
    (None, "SEGURO SOCIAL", TEXTO, 14),
    (None, "CURP", TEXTO, 21),
    (None, "RFC", TEXTO, 15),
    (None, "CORREO ELECTRONICO", TEXTO, 28),
    (None, "BENEFICIARIOS", TEXTO, 32),
    (None, "PUESTO", TEXTO, 24),
    (None, "FECHA DE INGRESO", FECHA, 14),
    (None, "FECHA ALTA IMSS", FECHA, 14),
    (None, "JORNADA DE TRABAJO (Horario)", TEXTO, 30),
    ("HORARIO LUNES A VIERNES", "ENTRADA", HORA, 11),
    ("HORARIO LUNES A VIERNES", "SALIDA", HORA, 11),
    (None, "HORARIO DE COMIDA", TEXTO, 18),
    ("HORARIO SÁBADO", "ENTRADA", HORA, 11),
    ("HORARIO SÁBADO", "SALIDA", HORA, 11),
    ("SALARIO DIARIO IMSS", "NÚMERO", IMPORTE, 12),
    ("SALARIO DIARIO IMSS", "LETRA", TEXTO, 38),
    (None, "DIA DE PAGO", TEXTO, 14),
    (None, "DIA DE DESCANSO", TEXTO, 16),
    (None, "VIGENCIA DEL CONTRATO", TEXTO, 20),
]

PATRON = [
    (None, "Razon Social", TEXTO, 36),
    (None, "Actividad de la empresa", TEXTO, 36),
    (None, "RFC", TEXTO, 15),
    ("DOMICILIO", "Numero", TEXTO, 9),
    ("DOMICILIO", "Calle", TEXTO, 22),
    ("DOMICILIO", "Colonia", TEXTO, 20),
    ("DOMICILIO", "Municipio", TEXTO, 20),
    ("DOMICILIO", "Codigo Postal", TEXTO, 12),
    (None, "Representante legal", TEXTO, 28),
    ("ACTA CONSTITUTIVA", "Numero", TEXTO, 10),
    ("ACTA CONSTITUTIVA", "Fecha", FECHA, 12),
    ("ACTA CONSTITUTIVA", "Nombre notario", TEXTO, 28),
    ("ACTA CONSTITUTIVA", "Numero notario", TEXTO, 10),
    ("ACTA CONSTITUTIVA", "Ciudad", TEXTO, 22),
]

INSTRUCCIONES = [
    "FORMATO DE CAPTURA — CONTRATOS INDIVIDUALES DE TRABAJO",
    "",
    "1. Capture un trabajador por renglón en la hoja «Trabajadores», a partir de la fila 3.",
    "2. Capture los datos de la empresa en la fila 3 de la hoja «Patron».",
    "3. Fechas: día/mes/año (por ejemplo 15/01/2026). Horas: formato de 24 horas (por ejemplo 09:00).",
    "4. HORARIO DE COMIDA: hora de inicio y de fin (por ejemplo 14:00 a 15:00).",
    "5. Si el trabajador no labora el sábado, deje vacías las columnas de HORARIO SÁBADO.",
    "6. SALARIO DIARIO IMSS: el número sin símbolo de pesos; en LETRA, el importe con letra "
    "terminado en /100 M.N. (por ejemplo SEISCIENTOS VEINTE PESOS 00/100 M.N.).",
    "7. Escriba el PUESTO igual en todos los trabajadores que lo compartan.",
    "8. No agregue ni cambie los encabezados de las filas 1 y 2.",
]

NAVY, DORADO, BLANCO = "15223B", "A8864F", "FFFFFF"
BORDE = Border(*(Side(style="thin", color="D9D2C3"),) * 4)


def hoja_de_captura(hoja, columnas, filas):
    for indice, (grupo, encabezado, formato, ancho) in enumerate(columnas, 1):
        letra = hoja.cell(2, indice).column_letter
        hoja.column_dimensions[letra].width = ancho
        celda = hoja.cell(2, indice, encabezado)
        celda.font = Font(name=FUENTE, bold=True, color=BLANCO, size=11)
        celda.fill = PatternFill("solid", fgColor=NAVY)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = BORDE
        for fila in range(3, 3 + filas):
            dato = hoja.cell(fila, indice)
            dato.number_format = formato
            dato.font = Font(name=FUENTE, size=11)
            dato.border = BORDE

    # Grupos (fila 1) en celdas combinadas.
    inicio = 1
    while inicio <= len(columnas):
        grupo = columnas[inicio - 1][0]
        fin = inicio
        while fin < len(columnas) and columnas[fin][0] == grupo:
            fin += 1
        if grupo:
            celda = hoja.cell(1, inicio, grupo)
            celda.font = Font(name=FUENTE, bold=True, color=BLANCO, size=11)
            celda.fill = PatternFill("solid", fgColor=DORADO)
            celda.alignment = Alignment(horizontal="center", vertical="center")
            if fin > inicio:
                hoja.merge_cells(start_row=1, start_column=inicio, end_row=1, end_column=fin)
        inicio = fin + 1

    hoja.row_dimensions[2].height = 32
    hoja.freeze_panes = "B3"


libro = Workbook()
hoja_de_captura(libro.active, TRABAJADORES, FILAS_DE_CAPTURA)
libro.active.title = "Trabajadores"
patron = libro.create_sheet("Patron")
hoja_de_captura(patron, PATRON, 1)

instrucciones = libro.create_sheet("Instrucciones")
instrucciones.column_dimensions["A"].width = 120
for fila, texto in enumerate(INSTRUCCIONES, 1):
    celda = instrucciones.cell(fila, 1, texto)
    celda.font = Font(name=FUENTE, size=12, bold=(fila == 1), color=NAVY if fila == 1 else "2A2E35")
    celda.alignment = Alignment(wrap_text=True, vertical="top")

libro.save(DESTINO)
print(f"Creado: {DESTINO.name}")
